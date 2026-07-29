"""
Finance Utility Suite
Universal Excel Reporting Engine

Author  : Shiva Kumar
Version : 1.37.2
"""

from __future__ import annotations

from abc import ABC, abstractmethod
from datetime import datetime
from pathlib import Path
from typing import Any, Final

import pandas as pd
from openpyxl.worksheet.worksheet import Worksheet


class ReportEngineError(RuntimeError):
    """Raised when a report cannot be generated safely."""


class ExcelReportEngine(ABC):
    """Reusable base class for Finance Utility Suite Excel reports.

    Subclasses provide report-specific workbook content through ``build``.
    The engine owns output validation, workbook lifecycle, metadata,
    common page setup, and final save/error handling.
    """

    ENGINE_VERSION: Final[str] = "1.37.2"
    APPLICATION_NAME: Final[str] = "Finance Utility Suite"
    DEFAULT_EXTENSION: Final[str] = ".xlsx"
    DEFAULT_ORIENTATION: Final[str] = "landscape"
    CONFIDENTIAL_FOOTER: Final[str] = (
        "Finance Utility Suite | Confidential | Page &P of &N"
    )

    def __init__(
        self,
        *,
        report_title: str,
        report_subject: str,
        report_description: str,
        report_keywords: str = "",
        creator: str | None = None,
    ) -> None:
        self.report_title = str(report_title).strip()
        self.report_subject = str(report_subject).strip()
        self.report_description = str(report_description).strip()
        self.report_keywords = str(report_keywords).strip()
        self.creator = str(creator or self.APPLICATION_NAME).strip()
        self.export_time: datetime | None = None

    def export(self, filepath: str | Path) -> Path:
        """Build and save a complete report workbook."""

        output_path = self._normalize_output_path(filepath)
        self.export_time = datetime.now()

        try:
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                placeholder = writer.book.create_sheet("_ReportBuild")
                self.build(writer, self.export_time)
                if len(writer.book.worksheets) > 1 and "_ReportBuild" in writer.book.sheetnames:
                    writer.book.remove(placeholder)
                self._apply_workbook_metadata(writer.book, self.export_time)
                self.finalize(writer)
        except Exception as exc:
            raise ReportEngineError(
                f"Unable to generate {self.report_title}: {exc}"
            ) from exc

        if not output_path.exists() or output_path.stat().st_size <= 0:
            raise ReportEngineError(
                f"The report file was not created correctly: {output_path}"
            )

        return output_path

    @abstractmethod
    def build(self, writer: pd.ExcelWriter, export_time: datetime) -> None:
        """Populate the workbook with report-specific worksheets."""

    def finalize(self, writer: pd.ExcelWriter) -> None:
        """Hook for report-specific final workbook adjustments."""

    def configure_worksheet(
        self,
        worksheet: Worksheet,
        *,
        orientation: str | None = None,
        fit_to_width: int = 1,
        fit_to_height: int | None = None,
        show_gridlines: bool = False,
        freeze_panes: str | None = None,
        print_area: str | None = None,
        repeat_rows: str | None = None,
        confidential_footer: bool = True,
    ) -> None:
        """Apply consistent Finance Utility Suite worksheet settings."""

        worksheet.sheet_view.showGridLines = show_gridlines
        worksheet.freeze_panes = freeze_panes
        worksheet.page_setup.orientation = orientation or self.DEFAULT_ORIENTATION
        worksheet.page_setup.fitToWidth = fit_to_width
        if fit_to_height is not None:
            worksheet.page_setup.fitToHeight = fit_to_height
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True

        if print_area:
            worksheet.print_area = print_area
        if repeat_rows:
            worksheet.print_title_rows = repeat_rows
        if confidential_footer:
            worksheet.oddFooter.center.text = self.CONFIDENTIAL_FOOTER

    @staticmethod
    def autosize_columns(
        worksheet: Worksheet,
        dataframe: pd.DataFrame,
        *,
        minimum_width: int = 10,
        maximum_width: int = 28,
        sample_rows: int = 500,
        padding: int = 3,
    ) -> None:
        """Set bounded column widths from dataframe content."""

        for index, column in enumerate(dataframe.columns, start=1):
            lengths = [len(str(column))]
            lengths.extend(
                len(str(value)) for value in dataframe[column].head(sample_rows)
            )
            width = max(minimum_width, min(max(lengths) + padding, maximum_width))
            letter = worksheet.cell(row=1, column=index).column_letter
            worksheet.column_dimensions[letter].width = width

    def _apply_workbook_metadata(
        self,
        workbook: Any,
        export_time: datetime,
    ) -> None:
        properties = workbook.properties
        properties.title = self.report_title
        properties.subject = self.report_subject
        properties.creator = self.creator
        properties.keywords = self.report_keywords
        properties.description = self.report_description
        properties.created = export_time
        properties.modified = export_time
        properties.lastModifiedBy = self.creator

    def _normalize_output_path(self, filepath: str | Path) -> Path:
        raw_path = Path(filepath).expanduser()
        if not str(raw_path).strip():
            raise ReportEngineError("An output file path is required.")

        if raw_path.suffix.lower() != self.DEFAULT_EXTENSION:
            raw_path = raw_path.with_suffix(self.DEFAULT_EXTENSION)

        return raw_path.resolve()
