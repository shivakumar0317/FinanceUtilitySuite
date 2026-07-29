"""Shared Excel styles for Finance Utility Suite reports."""

from __future__ import annotations

from dataclasses import dataclass

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


@dataclass(frozen=True)
class ExcelTheme:
    dark_blue: str = "17365D"
    medium_blue: str = "2F75B5"
    light_blue: str = "D9EAF7"
    primary: str = "17365D"
    secondary: str = "D9EAF7"
    pale_blue: str = "EAF3F8"
    white: str = "FFFFFF"
    dark_text: str = "1F2937"
    muted_text: str = "64748B"
    border: str = "CBD5E1"
    green: str = "16A34A"
    amber: str = "CA8A04"
    orange: str = "EA580C"
    red: str = "DC2626"


THEME = ExcelTheme()


def thin_border() -> Border:
    side = Side(style="thin", color=THEME.border)
    return Border(left=side, right=side, top=side, bottom=side)


def solid_fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


def title_font(size: int = 18) -> Font:
    return Font(name="Segoe UI", size=size, bold=True, color=THEME.white)


def section_font(size: int = 11) -> Font:
    return Font(name="Segoe UI", size=size, bold=True, color=THEME.white)


def label_font(size: int = 9) -> Font:
    return Font(name="Segoe UI", size=size, bold=True, color=THEME.muted_text)


def value_font(size: int = 14, color: str | None = None) -> Font:
    return Font(
        name="Segoe UI",
        size=size,
        bold=True,
        color=color or THEME.dark_text,
    )


def normal_font(size: int = 10, color: str | None = None) -> Font:
    return Font(name="Segoe UI", size=size, color=color or THEME.dark_text)


def centered() -> Alignment:
    return Alignment(horizontal="center", vertical="center")
