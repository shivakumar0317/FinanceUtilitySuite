"""
Finance Utility Suite
Executive Recommendation Engine

Author  : Shiva Kumar
Version : 1.35.0
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from desktop.reports.excel_styles import (
    THEME,
    centered,
    normal_font,
    section_font,
    solid_fill,
    thin_border,
    title_font,
    value_font,
)


@dataclass(frozen=True)
class Recommendation:
    """One explainable portfolio recommendation."""

    rule_id: str
    area: str
    severity: str
    status: str
    message: str
    action: str
    review_period: str


class ExecutiveRecommendationEngine:
    """Evaluate portfolio metrics and build an executive intelligence sheet."""

    VERSION = "1.35.0"

    SEVERITY_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3, "Info": 4}

    def __init__(
        self,
        account_id: str,
        holdings_df: pd.DataFrame,
        metrics: dict[str, float],
        summary_data: dict[str, Any] | None = None,
    ) -> None:
        self.account_id = str(account_id).strip()
        self.holdings_df = holdings_df.copy()
        self.metrics = dict(metrics)
        self.summary_data = dict(summary_data or {})

    def build(self, worksheet: Worksheet) -> None:
        """Populate a management-ready Executive Intelligence worksheet."""

        recommendations = self.evaluate()
        worksheet.sheet_view.showGridLines = False
        self._write_title(worksheet)
        self._write_scorecard(worksheet, recommendations)
        self._write_risk_matrix(worksheet)
        self._write_recommendations(worksheet, recommendations)
        self._write_action_plan(worksheet, recommendations)
        self._apply_page_setup(worksheet)

    def evaluate(self) -> list[Recommendation]:
        """Return deterministic, explainable recommendations."""

        recs: list[Recommendation] = []
        largest = float(self.metrics.get("largest_holding", 0.0))
        top_three = float(self.metrics.get("top_three", 0.0))
        diversification = float(self.metrics.get("diversification", 0.0))
        total_mtm = float(self.metrics.get("total_mtm", 0.0))
        total_exposure = float(self.metrics.get("total_exposure", 0.0))
        total_margin = float(self.metrics.get("total_margin", 0.0))
        holding_count = int(self.metrics.get("holding_count", 0.0))

        if largest >= 35:
            recs.append(self._rec("CONC-01", "Largest Holding", "Critical", "Critical", "Largest holding is above 35% of portfolio exposure.", "Review and reduce single-stock concentration subject to approved risk policy.", "Immediate"))
        elif largest >= 20:
            recs.append(self._rec("CONC-02", "Largest Holding", "High", "High", "Largest holding exceeds the preferred 20% concentration level.", "Review concentration and consider gradual rebalancing.", "This Week"))
        elif largest >= 15:
            recs.append(self._rec("CONC-03", "Largest Holding", "Medium", "Monitor", "Largest holding is approaching the preferred concentration limit.", "Monitor the position and avoid further concentration without approval.", "This Month"))
        else:
            recs.append(self._rec("CONC-04", "Largest Holding", "Info", "Healthy", "Largest holding remains below 15%.", "Continue periodic concentration monitoring.", "Quarterly"))

        if top_three >= 60:
            recs.append(self._rec("TOP3-01", "Top 3 Holdings", "Critical", "Critical", "Top three holdings represent at least 60% of portfolio exposure.", "Prioritize a concentration review of the top three positions.", "Immediate"))
        elif top_three >= 50:
            recs.append(self._rec("TOP3-02", "Top 3 Holdings", "High", "High", "Top three holdings represent at least half of portfolio exposure.", "Assess whether exposure should be distributed across more securities.", "This Week"))
        elif top_three >= 40:
            recs.append(self._rec("TOP3-03", "Top 3 Holdings", "Medium", "Monitor", "Top three holdings account for a meaningful share of exposure.", "Review the top three holdings during the monthly portfolio review.", "This Month"))
        else:
            recs.append(self._rec("TOP3-04", "Top 3 Holdings", "Info", "Healthy", "Top three holdings are within a balanced range.", "Maintain routine monitoring.", "Quarterly"))

        if diversification < 40:
            recs.append(self._rec("DIV-01", "Diversification", "Critical", "Critical", "Diversification score is below 40.", "Increase diversification across suitable securities and sectors.", "Immediate"))
        elif diversification < 60:
            recs.append(self._rec("DIV-02", "Diversification", "High", "Weak", "Diversification score is below 60.", "Review opportunities to reduce portfolio concentration.", "This Week"))
        elif diversification < 75:
            recs.append(self._rec("DIV-03", "Diversification", "Medium", "Moderate", "Diversification is adequate but can be improved.", "Review portfolio balance during the next monthly review.", "This Month"))
        else:
            recs.append(self._rec("DIV-04", "Diversification", "Info", "Excellent", "Portfolio diversification is healthy.", "Continue periodic diversification checks.", "Quarterly"))

        mtm_ratio = (total_mtm / total_exposure * 100) if total_exposure else 0.0
        if mtm_ratio <= -15:
            recs.append(self._rec("MTM-01", "Mark to Market", "Critical", "Critical", "Portfolio MTM loss exceeds 15% of exposure.", "Review loss-making positions and apply approved risk controls immediately.", "Immediate"))
        elif mtm_ratio <= -8:
            recs.append(self._rec("MTM-02", "Mark to Market", "High", "High", "Portfolio MTM loss exceeds 8% of exposure.", "Review the largest negative MTM contributors.", "This Week"))
        elif total_mtm < 0:
            recs.append(self._rec("MTM-03", "Mark to Market", "Medium", "Monitor", "Portfolio currently has a negative aggregate MTM.", "Monitor negative MTM positions and reassess risk limits.", "This Week"))
        else:
            recs.append(self._rec("MTM-04", "Mark to Market", "Info", "Healthy", "Aggregate MTM is non-negative.", "Continue regular MTM monitoring.", "Weekly"))

        margin_ratio = (total_margin / total_exposure * 100) if total_exposure else 0.0
        if margin_ratio >= 60:
            recs.append(self._rec("MAR-01", "MTF Margin", "High", "High", "MTF margin exceeds 60% of total exposure.", "Review funding efficiency and margin requirements.", "This Week"))
        elif margin_ratio >= 40:
            recs.append(self._rec("MAR-02", "MTF Margin", "Medium", "Monitor", "MTF margin is elevated relative to exposure.", "Monitor funding exposure and collateral adequacy.", "This Month"))
        else:
            recs.append(self._rec("MAR-03", "MTF Margin", "Info", "Healthy", "MTF margin is within the configured monitoring range.", "Continue routine funding review.", "Monthly"))

        if holding_count <= 1:
            recs.append(self._rec("COUNT-01", "Holdings Count", "Critical", "Critical", "Portfolio contains one or fewer holdings.", "Add suitable holdings to reduce single-security dependency.", "Immediate"))
        elif holding_count < 5:
            recs.append(self._rec("COUNT-02", "Holdings Count", "High", "Weak", "Portfolio contains fewer than five holdings.", "Review diversification across additional suitable securities.", "This Week"))
        elif holding_count < 10:
            recs.append(self._rec("COUNT-03", "Holdings Count", "Medium", "Moderate", "Portfolio has a limited number of holdings.", "Assess whether additional diversification is appropriate.", "This Month"))
        else:
            recs.append(self._rec("COUNT-04", "Holdings Count", "Info", "Healthy", "Portfolio has a broad holdings base.", "Maintain periodic review of position quality.", "Quarterly"))

        return sorted(recs, key=lambda item: (self.SEVERITY_ORDER.get(item.severity, 99), item.area))

    @staticmethod
    def _rec(rule_id: str, area: str, severity: str, status: str, message: str, action: str, review_period: str) -> Recommendation:
        return Recommendation(rule_id, area, severity, status, message, action, review_period)

    def _write_title(self, worksheet: Worksheet) -> None:
        worksheet.merge_cells("A1:L1")
        worksheet["A1"] = "EXECUTIVE INTELLIGENCE"
        worksheet["A1"].fill = solid_fill(THEME.dark_blue)
        worksheet["A1"].font = title_font(18)
        worksheet["A1"].alignment = centered()
        worksheet.row_dimensions[1].height = 30

        worksheet.merge_cells("A2:L2")
        worksheet["A2"] = f"Client Account: {self.account_id} | Rule-based decision support"
        worksheet["A2"].fill = solid_fill(THEME.medium_blue)
        worksheet["A2"].font = title_font(11)
        worksheet["A2"].alignment = centered()

    def _write_scorecard(self, worksheet: Worksheet, recommendations: list[Recommendation]) -> None:
        score = self._health_score()
        label = self._health_label(score)
        highest = recommendations[0].severity if recommendations else "Info"

        worksheet.merge_cells("A4:H4")
        worksheet["A4"] = "EXECUTIVE SCORECARD"
        worksheet["A4"].fill = solid_fill(THEME.dark_blue)
        worksheet["A4"].font = section_font()

        cards = [
            ("Portfolio Health", score, '0.0" / 100"', self._health_color(score)),
            ("Overall Rating", label, "General", self._health_color(score)),
            ("Diversification", self._status_diversification(), "General", self._status_color(self._status_diversification())),
            ("Concentration", self._status_concentration(), "General", self._status_color(self._status_concentration())),
            ("MTM", self._status_mtm(), "General", self._status_color(self._status_mtm())),
            ("Highest Alert", highest, "General", self._severity_color(highest)),
        ]
        positions = [(6, 1), (6, 4), (6, 7), (9, 1), (9, 4), (9, 7)]
        border = thin_border()
        for (label_text, value, number_format, color), (row, col) in zip(cards, positions):
            worksheet.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 2)
            worksheet.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 2)
            label_cell = worksheet.cell(row, col, label_text)
            value_cell = worksheet.cell(row + 1, col, value)
            label_cell.fill = solid_fill(THEME.pale_blue)
            label_cell.font = Font(name="Segoe UI", size=9, bold=True, color=THEME.muted_text)
            label_cell.alignment = centered()
            value_cell.fill = solid_fill(THEME.white)
            value_cell.font = value_font(14, color)
            value_cell.alignment = centered()
            value_cell.number_format = number_format
            for rr in (row, row + 1):
                for cc in range(col, col + 3):
                    worksheet.cell(rr, cc).border = border

    def _write_risk_matrix(self, worksheet: Worksheet) -> None:
        row = 13
        worksheet.merge_cells(f"A{row}:L{row}")
        worksheet[f"A{row}"] = "RISK MATRIX"
        worksheet[f"A{row}"].fill = solid_fill(THEME.dark_blue)
        worksheet[f"A{row}"].font = section_font()

        headers = ["Area", "Status", "Comment"]
        for col, header in enumerate(headers, start=1):
            cell = worksheet.cell(row + 1, col, header)
            cell.fill = solid_fill(THEME.medium_blue)
            cell.font = Font(name="Segoe UI", size=10, bold=True, color=THEME.white)
            cell.alignment = centered()
            cell.border = thin_border()
        worksheet.merge_cells(start_row=row + 1, start_column=3, end_row=row + 1, end_column=12)

        rows = [
            ("Diversification", self._status_diversification(), self._comment_diversification()),
            ("Concentration", self._status_concentration(), self._comment_concentration()),
            ("Mark to Market", self._status_mtm(), self._comment_mtm()),
            ("MTF Margin", self._status_margin(), self._comment_margin()),
            ("Largest Holding", self._status_largest(), self._comment_largest()),
        ]
        for offset, (area, status, comment) in enumerate(rows, start=2):
            target = row + offset
            worksheet.cell(target, 1, area)
            worksheet.cell(target, 2, status)
            worksheet.merge_cells(start_row=target, start_column=3, end_row=target, end_column=12)
            worksheet.cell(target, 3, comment)
            for col in range(1, 13):
                cell = worksheet.cell(target, col)
                cell.border = thin_border()
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.fill = solid_fill(THEME.white if offset % 2 == 0 else THEME.pale_blue)
            worksheet.cell(target, 2).font = Font(name="Segoe UI", size=10, bold=True, color=self._status_color(status))

    def _write_recommendations(self, worksheet: Worksheet, recommendations: list[Recommendation]) -> None:
        start = 22
        worksheet.merge_cells(f"A{start}:L{start}")
        worksheet[f"A{start}"] = "RECOMMENDATIONS"
        worksheet[f"A{start}"].fill = solid_fill(THEME.dark_blue)
        worksheet[f"A{start}"].font = section_font()

        headers = ["Priority", "Area", "Finding", "Recommended Action", "Review"]
        columns = [(1, 1), (2, 3), (4, 7), (8, 11), (12, 12)]
        for header, (start_col, end_col) in zip(headers, columns):
            if end_col > start_col:
                worksheet.merge_cells(start_row=start + 1, start_column=start_col, end_row=start + 1, end_column=end_col)
            cell = worksheet.cell(start + 1, start_col, header)
            cell.fill = solid_fill(THEME.medium_blue)
            cell.font = Font(name="Segoe UI", size=10, bold=True, color=THEME.white)
            cell.alignment = centered()
            for col in range(start_col, end_col + 1):
                worksheet.cell(start + 1, col).border = thin_border()

        for index, rec in enumerate(recommendations, start=1):
            row = start + 1 + index
            values = [rec.severity, rec.area, rec.message, rec.action, rec.review_period]
            for value, (start_col, end_col) in zip(values, columns):
                if end_col > start_col:
                    worksheet.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
                cell = worksheet.cell(row, start_col, value)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.font = normal_font(9)
                for col in range(start_col, end_col + 1):
                    worksheet.cell(row, col).border = thin_border()
                    worksheet.cell(row, col).fill = solid_fill(THEME.white if index % 2 else THEME.pale_blue)
            worksheet.cell(row, 1).font = Font(name="Segoe UI", size=9, bold=True, color=self._severity_color(rec.severity))
            worksheet.row_dimensions[row].height = 42

    def _write_action_plan(self, worksheet: Worksheet, recommendations: list[Recommendation]) -> None:
        start = 32
        worksheet.merge_cells(f"A{start}:L{start}")
        worksheet[f"A{start}"] = "ACTION PLAN"
        worksheet[f"A{start}"].fill = solid_fill(THEME.dark_blue)
        worksheet[f"A{start}"].font = section_font()

        groups = ["Immediate", "This Week", "This Month", "Monthly", "Quarterly", "Weekly"]
        row = start + 2
        for period in groups:
            items = [rec for rec in recommendations if rec.review_period == period]
            if not items:
                continue
            worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
            worksheet.cell(row, 1, period.upper())
            worksheet.cell(row, 1).fill = solid_fill(THEME.light_blue)
            worksheet.cell(row, 1).font = Font(name="Segoe UI", size=10, bold=True, color=THEME.dark_text)
            worksheet.cell(row, 1).alignment = centered()
            for col in range(1, 4):
                worksheet.cell(row, col).border = thin_border()
            row += 1
            for item in items:
                worksheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
                worksheet.cell(row, 1, f"□ {item.action}")
                worksheet.cell(row, 1).alignment = Alignment(vertical="center", wrap_text=True)
                worksheet.cell(row, 1).font = normal_font(9)
                for col in range(1, 13):
                    worksheet.cell(row, col).border = thin_border()
                worksheet.row_dimensions[row].height = 28
                row += 1
            row += 1

    def _health_score(self) -> float:
        diversification = float(self.metrics.get("diversification", 0.0))
        concentration_component = max(0.0, 100.0 - float(self.metrics.get("concentration", 0.0)))
        largest = float(self.metrics.get("largest_holding", 0.0))
        largest_component = max(0.0, min(100.0, 100.0 - (largest * 2.5)))
        exposure = float(self.metrics.get("total_exposure", 0.0))
        mtm = float(self.metrics.get("total_mtm", 0.0))
        mtm_ratio = (mtm / exposure * 100) if exposure else 0.0
        mtm_component = max(0.0, min(100.0, 75.0 + (mtm_ratio * 2.0)))
        return max(0.0, min(100.0, diversification * 0.40 + concentration_component * 0.20 + largest_component * 0.20 + mtm_component * 0.20))

    @staticmethod
    def _health_label(score: float) -> str:
        if score >= 85: return "Excellent"
        if score >= 70: return "Good"
        if score >= 55: return "Moderate"
        if score >= 40: return "Weak"
        return "Critical"

    @staticmethod
    def _health_color(score: float) -> str:
        if score >= 85: return THEME.green
        if score >= 70: return "65A30D"
        if score >= 55: return THEME.amber
        if score >= 40: return THEME.orange
        return THEME.red

    def _status_diversification(self) -> str:
        value = float(self.metrics.get("diversification", 0.0))
        return "Excellent" if value >= 75 else "Moderate" if value >= 60 else "Weak" if value >= 40 else "Critical"

    def _status_concentration(self) -> str:
        value = float(self.metrics.get("largest_holding", 0.0))
        return "Critical" if value >= 35 else "High" if value >= 20 else "Monitor" if value >= 15 else "Healthy"

    def _status_mtm(self) -> str:
        exposure = float(self.metrics.get("total_exposure", 0.0))
        mtm = float(self.metrics.get("total_mtm", 0.0))
        ratio = mtm / exposure * 100 if exposure else 0.0
        return "Critical" if ratio <= -15 else "High" if ratio <= -8 else "Monitor" if mtm < 0 else "Healthy"

    def _status_margin(self) -> str:
        exposure = float(self.metrics.get("total_exposure", 0.0))
        margin = float(self.metrics.get("total_margin", 0.0))
        ratio = margin / exposure * 100 if exposure else 0.0
        return "High" if ratio >= 60 else "Monitor" if ratio >= 40 else "Healthy"

    def _status_largest(self) -> str:
        return self._status_concentration()

    def _comment_diversification(self) -> str:
        return f"Diversification score is {float(self.metrics.get('diversification', 0.0)):.1f}/100."

    def _comment_concentration(self) -> str:
        return f"Largest holding is {float(self.metrics.get('largest_holding', 0.0)):.2f}% and top three holdings are {float(self.metrics.get('top_three', 0.0)):.2f}%."

    def _comment_mtm(self) -> str:
        return f"Aggregate MTM is ₹{float(self.metrics.get('total_mtm', 0.0)):,.2f}."

    def _comment_margin(self) -> str:
        return f"Total MTF margin is ₹{float(self.metrics.get('total_margin', 0.0)):,.2f}."

    def _comment_largest(self) -> str:
        return f"Largest position represents {float(self.metrics.get('largest_holding', 0.0)):.2f}% of total exposure."

    @staticmethod
    def _severity_color(severity: str) -> str:
        return {"Critical": THEME.red, "High": THEME.orange, "Medium": THEME.amber, "Low": "65A30D", "Info": THEME.green}.get(severity, THEME.muted_text)

    @staticmethod
    def _status_color(status: str) -> str:
        return {"Critical": THEME.red, "High": THEME.orange, "Weak": THEME.orange, "Monitor": THEME.amber, "Moderate": THEME.amber, "Healthy": THEME.green, "Excellent": THEME.green}.get(status, THEME.muted_text)

    @staticmethod
    def _apply_page_setup(worksheet: Worksheet) -> None:
        for letter, width in {"A": 14, "B": 14, "C": 14, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15, "I": 15, "J": 15, "K": 15, "L": 16}.items():
            worksheet.column_dimensions[letter].width = width
        worksheet.freeze_panes = "A4"
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.oddFooter.center.text = "Finance Utility Suite | Confidential | Page &P of &N"
