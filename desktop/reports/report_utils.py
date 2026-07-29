"""
Finance Utility Suite
Reusable Excel Report Utilities

Shared worksheet formatting, table creation, number formats,
profit/loss highlighting, and report layout helpers.

Author  : Shiva Kumar
Version : 1.37.1
"""

from __future__ import annotations

from collections.abc import Iterable, Sequence
from datetime import date, datetime
from numbers import Number
from typing import Any

import pandas as pd
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from desktop.reports.excel_styles import (
    THEME,
    centered,
    label_font,
    normal_font,
    section_font,
    solid_fill,
    thin_border,
    title_font,
    value_font,
)


# ---------------------------------------------------------------------------
# Excel number formats
# ---------------------------------------------------------------------------

CURRENCY_FORMAT = '₹#,##0.00;[Red]-₹#,##0.00'
CURRENCY_FORMAT_NO_DECIMALS = '₹#,##0;[Red]-₹#,##0'
NUMBER_FORMAT = '#,##0.00;[Red]-#,##0.00'
NUMBER_FORMAT_NO_DECIMALS = '#,##0;[Red]-#,##0'
PERCENT_FORMAT = '0.00%;[Red]-0.00%'
PERCENT_VALUE_FORMAT = '0.00%;[Red]-0.00%'
DATE_FORMAT = 'dd-mmm-yyyy'
DATETIME_FORMAT = 'dd-mmm-yyyy hh:mm AM/PM'

POSITIVE_FILL = PatternFill(
    fill_type="solid",
    fgColor="E2F0D9",
)

NEGATIVE_FILL = PatternFill(
    fill_type="solid",
    fgColor="FCE4D6",
)

NEUTRAL_FILL = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)

POSITIVE_FONT = Font(
    color="008000",
    bold=True,
)

NEGATIVE_FONT = Font(
    color="C00000",
    bold=True,
)

NEUTRAL_FONT = Font(
    color="7F6000",
    bold=True,
)


# ---------------------------------------------------------------------------
# Value formatting helpers
# ---------------------------------------------------------------------------

def safe_float(value: Any, default: float = 0.0) -> float:
    """Convert a value to float without raising an exception."""

    if value is None:
        return default

    try:
        if pd.isna(value):
            return default
    except (TypeError, ValueError):
        pass

    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def format_currency(
    value: Any,
    *,
    decimals: int = 2,
    symbol: str = "₹",
) -> str:
    """Return a display-friendly currency string."""

    number = safe_float(value)

    if decimals <= 0:
        return f"{symbol}{number:,.0f}"

    return f"{symbol}{number:,.{decimals}f}"


def format_number(
    value: Any,
    *,
    decimals: int = 2,
) -> str:
    """Return a display-friendly numeric string."""

    number = safe_float(value)

    if decimals <= 0:
        return f"{number:,.0f}"

    return f"{number:,.{decimals}f}"


def format_percent(
    value: Any,
    *,
    decimals: int = 2,
    value_is_fraction: bool = False,
) -> str:
    """Return a display-friendly percentage string.

    Examples:
        format_percent(12.5) -> "12.50%"
        format_percent(0.125, value_is_fraction=True) -> "12.50%"
    """

    number = safe_float(value)

    if value_is_fraction:
        number *= 100

    return f"{number:,.{decimals}f}%"


def normalize_excel_value(value: Any) -> Any:
    """Convert pandas and NumPy values into Excel-safe Python values."""

    if value is None:
        return None

    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()

    if isinstance(value, datetime):
        return value

    if isinstance(value, date):
        return value

    if hasattr(value, "item"):
        try:
            return value.item()
        except (AttributeError, ValueError):
            pass

    return value


# ---------------------------------------------------------------------------
# Worksheet title and section helpers
# ---------------------------------------------------------------------------

def write_report_title(
    worksheet: Worksheet,
    *,
    title: str,
    subtitle: str | None = None,
    export_time: datetime | None = None,
    start_row: int = 1,
    start_column: int = 1,
    end_column: int = 8,
) -> int:
    """Write the report heading and return the next available row."""

    title_row = start_row
    subtitle_row = title_row + 1
    timestamp_row = subtitle_row + 1

    worksheet.merge_cells(
        start_row=title_row,
        start_column=start_column,
        end_row=title_row,
        end_column=end_column,
    )

    title_cell = worksheet.cell(
        row=title_row,
        column=start_column,
        value=title,
    )
    title_cell.font = title_font()
    title_cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[title_row].height = 28

    current_row = subtitle_row

    if subtitle:
        worksheet.merge_cells(
            start_row=subtitle_row,
            start_column=start_column,
            end_row=subtitle_row,
            end_column=end_column,
        )

        subtitle_cell = worksheet.cell(
            row=subtitle_row,
            column=start_column,
            value=subtitle,
        )
        subtitle_cell.font = normal_font()
        subtitle_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        current_row = timestamp_row

    if export_time:
        worksheet.merge_cells(
            start_row=current_row,
            start_column=start_column,
            end_row=current_row,
            end_column=end_column,
        )

        timestamp_cell = worksheet.cell(
            row=current_row,
            column=start_column,
            value=f"Generated on: {export_time:%d-%b-%Y %I:%M %p}",
        )
        timestamp_cell.font = normal_font()
        timestamp_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        current_row += 1

    return current_row + 1


def write_section_header(
    worksheet: Worksheet,
    *,
    row: int,
    title: str,
    start_column: int = 1,
    end_column: int = 8,
) -> int:
    """Write a merged section heading and return the next row."""

    worksheet.merge_cells(
        start_row=row,
        start_column=start_column,
        end_row=row,
        end_column=end_column,
    )

    cell = worksheet.cell(
        row=row,
        column=start_column,
        value=title,
    )
    cell.font = section_font()
    cell.fill = solid_fill(THEME.primary)
    cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[row].height = 22

    return row + 1


# ---------------------------------------------------------------------------
# KPI and summary helpers
# ---------------------------------------------------------------------------

def write_kpi_cards(
    worksheet: Worksheet,
    *,
    metrics: Sequence[tuple[str, Any]],
    start_row: int,
    start_column: int = 1,
    cards_per_row: int = 4,
    card_width: int = 2,
    number_formats: dict[str, str] | None = None,
) -> int:
    """Write KPI cards and return the first row after the card area.

    Each metric must be supplied as:

        ("Portfolio Value", 1500000)

    ``number_formats`` may map a metric label to an Excel number format.
    """

    if not metrics:
        return start_row

    formats = number_formats or {}

    for index, (label, value) in enumerate(metrics):
        card_row = start_row + ((index // cards_per_row) * 3)
        card_column = start_column + (
            (index % cards_per_row) * card_width
        )

        end_column = card_column + card_width - 1

        worksheet.merge_cells(
            start_row=card_row,
            start_column=card_column,
            end_row=card_row,
            end_column=end_column,
        )

        worksheet.merge_cells(
            start_row=card_row + 1,
            start_column=card_column,
            end_row=card_row + 1,
            end_column=end_column,
        )

        label_cell = worksheet.cell(
            row=card_row,
            column=card_column,
            value=label,
        )
        label_cell.font = label_font()
        label_cell.fill = solid_fill(THEME.secondary)
        label_cell.alignment = centered()
        label_cell.border = thin_border()

        value_cell = worksheet.cell(
            row=card_row + 1,
            column=card_column,
            value=normalize_excel_value(value),
        )
        value_cell.font = value_font()
        value_cell.alignment = centered()
        value_cell.border = thin_border()

        if label in formats:
            value_cell.number_format = formats[label]

        for column in range(card_column, end_column + 1):
            worksheet.cell(
                row=card_row,
                column=column,
            ).fill = solid_fill(THEME.secondary)

            worksheet.cell(
                row=card_row,
                column=column,
            ).border = thin_border()

            worksheet.cell(
                row=card_row + 1,
                column=column,
            ).border = thin_border()

        worksheet.row_dimensions[card_row].height = 20
        worksheet.row_dimensions[card_row + 1].height = 28

    total_card_rows = ((len(metrics) - 1) // cards_per_row) + 1
    return start_row + (total_card_rows * 3)


def write_key_value_summary(
    worksheet: Worksheet,
    *,
    values: Sequence[tuple[str, Any]],
    start_row: int,
    start_column: int = 1,
    value_column: int | None = None,
    number_formats: dict[str, str] | None = None,
) -> int:
    """Write a vertical key-value summary and return the next row."""

    formats = number_formats or {}
    value_col = value_column or start_column + 1

    for offset, (label, value) in enumerate(values):
        row = start_row + offset

        label_cell = worksheet.cell(
            row=row,
            column=start_column,
            value=label,
        )
        label_cell.font = label_font()
        label_cell.fill = solid_fill(THEME.secondary)
        label_cell.border = thin_border()
        label_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        value_cell = worksheet.cell(
            row=row,
            column=value_col,
            value=normalize_excel_value(value),
        )
        value_cell.font = value_font()
        value_cell.border = thin_border()
        value_cell.alignment = Alignment(
            horizontal="right",
            vertical="center",
        )

        if label in formats:
            value_cell.number_format = formats[label]

    return start_row + len(values)


# ---------------------------------------------------------------------------
# DataFrame writing and table helpers
# ---------------------------------------------------------------------------

def write_dataframe(
    worksheet: Worksheet,
    dataframe: pd.DataFrame,
    *,
    start_row: int = 1,
    start_column: int = 1,
    include_index: bool = False,
    header_fill: str | None = None,
) -> tuple[int, int]:
    """Write a DataFrame and return its ending row and column."""

    if dataframe is None:
        raise ValueError("A dataframe is required.")

    frame = dataframe.copy()

    if include_index:
        frame = frame.reset_index()

    if frame.columns.empty:
        return start_row, start_column

    fill_color = header_fill or THEME.primary

    for column_offset, column_name in enumerate(frame.columns):
        cell = worksheet.cell(
            row=start_row,
            column=start_column + column_offset,
            value=str(column_name),
        )
        cell.font = section_font()
        cell.fill = solid_fill(fill_color)
        cell.alignment = centered()
        cell.border = thin_border()

    for row_offset, row_values in enumerate(
        frame.itertuples(index=False, name=None),
        start=1,
    ):
        for column_offset, value in enumerate(row_values):
            cell = worksheet.cell(
                row=start_row + row_offset,
                column=start_column + column_offset,
                value=normalize_excel_value(value),
            )
            cell.font = normal_font()
            cell.border = thin_border()
            cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
            )

            if isinstance(value, (datetime, pd.Timestamp)):
                cell.number_format = DATETIME_FORMAT
            elif isinstance(value, date):
                cell.number_format = DATE_FORMAT

    end_row = start_row + len(frame)
    end_column = start_column + len(frame.columns) - 1

    return end_row, end_column


def create_excel_table(
    worksheet: Worksheet,
    *,
    start_row: int,
    start_column: int,
    end_row: int,
    end_column: int,
    table_name: str,
    style_name: str = "TableStyleMedium2",
) -> Table | None:
    """Create an Excel table for a worksheet range."""

    if end_row <= start_row:
        return None

    if end_column < start_column:
        return None

    reference = (
        f"{get_column_letter(start_column)}{start_row}:"
        f"{get_column_letter(end_column)}{end_row}"
    )

    safe_name = sanitize_table_name(table_name)

    existing_names = {
        table.displayName
        for sheet in worksheet.parent.worksheets
        for table in sheet.tables.values()
    }

    original_name = safe_name
    suffix = 1

    while safe_name in existing_names:
        safe_name = f"{original_name}_{suffix}"
        suffix += 1

    table = Table(
        displayName=safe_name,
        ref=reference,
    )

    table.tableStyleInfo = TableStyleInfo(
        name=style_name,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    worksheet.add_table(table)

    return table


def sanitize_table_name(name: str) -> str:
    """Return an Excel-compatible table name."""

    cleaned = "".join(
        character if character.isalnum() or character == "_" else "_"
        for character in str(name).strip()
    )

    if not cleaned:
        cleaned = "ReportTable"

    if cleaned[0].isdigit():
        cleaned = f"T_{cleaned}"

    return cleaned[:255]


# ---------------------------------------------------------------------------
# Column formatting
# ---------------------------------------------------------------------------

def apply_number_format(
    worksheet: Worksheet,
    *,
    column: int | str,
    start_row: int,
    end_row: int,
    number_format: str,
) -> None:
    """Apply an Excel number format to a worksheet column range."""

    column_index = (
        column
        if isinstance(column, int)
        else worksheet[column][0].column
    )

    for row in range(start_row, end_row + 1):
        worksheet.cell(
            row=row,
            column=column_index,
        ).number_format = number_format


def apply_dataframe_number_formats(
    worksheet: Worksheet,
    dataframe: pd.DataFrame,
    *,
    start_row: int,
    start_column: int = 1,
    currency_columns: Iterable[str] = (),
    percentage_columns: Iterable[str] = (),
    numeric_columns: Iterable[str] = (),
    date_columns: Iterable[str] = (),
    percentage_values_are_fraction: bool = False,
) -> None:
    """Apply number formats using DataFrame column names."""

    if dataframe is None or dataframe.empty:
        return

    column_lookup = {
        str(column): start_column + index
        for index, column in enumerate(dataframe.columns)
    }

    first_data_row = start_row + 1
    last_data_row = start_row + len(dataframe)

    for column_name in currency_columns:
        column_index = column_lookup.get(str(column_name))
        if column_index is not None:
            apply_number_format(
                worksheet,
                column=column_index,
                start_row=first_data_row,
                end_row=last_data_row,
                number_format=CURRENCY_FORMAT,
            )

    percent_format = (
        PERCENT_FORMAT
        if percentage_values_are_fraction
        else '0.00"%"'
    )

    for column_name in percentage_columns:
        column_index = column_lookup.get(str(column_name))
        if column_index is not None:
            apply_number_format(
                worksheet,
                column=column_index,
                start_row=first_data_row,
                end_row=last_data_row,
                number_format=percent_format,
            )

    for column_name in numeric_columns:
        column_index = column_lookup.get(str(column_name))
        if column_index is not None:
            apply_number_format(
                worksheet,
                column=column_index,
                start_row=first_data_row,
                end_row=last_data_row,
                number_format=NUMBER_FORMAT,
            )

    for column_name in date_columns:
        column_index = column_lookup.get(str(column_name))
        if column_index is not None:
            apply_number_format(
                worksheet,
                column=column_index,
                start_row=first_data_row,
                end_row=last_data_row,
                number_format=DATE_FORMAT,
            )


# ---------------------------------------------------------------------------
# Profit/loss and conditional formatting
# ---------------------------------------------------------------------------

def apply_profit_loss_format(
    worksheet: Worksheet,
    *,
    column: int | str,
    start_row: int,
    end_row: int,
    zero_is_neutral: bool = True,
) -> None:
    """Apply static positive, negative, and neutral styling."""

    column_index = (
        column
        if isinstance(column, int)
        else worksheet[column][0].column
    )

    for row in range(start_row, end_row + 1):
        cell = worksheet.cell(
            row=row,
            column=column_index,
        )

        if not isinstance(cell.value, Number):
            continue

        if cell.value > 0:
            cell.font = POSITIVE_FONT
            cell.fill = POSITIVE_FILL
        elif cell.value < 0:
            cell.font = NEGATIVE_FONT
            cell.fill = NEGATIVE_FILL
        elif zero_is_neutral:
            cell.font = NEUTRAL_FONT
            cell.fill = NEUTRAL_FILL


def add_profit_loss_conditional_formatting(
    worksheet: Worksheet,
    *,
    column: int | str,
    start_row: int,
    end_row: int,
) -> None:
    """Add Excel-native conditional formatting for profit/loss values."""

    column_letter = (
        get_column_letter(column)
        if isinstance(column, int)
        else column
    )

    cell_range = (
        f"{column_letter}{start_row}:"
        f"{column_letter}{end_row}"
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="greaterThan",
            formula=["0"],
            fill=POSITIVE_FILL,
            font=POSITIVE_FONT,
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="lessThan",
            formula=["0"],
            fill=NEGATIVE_FILL,
            font=NEGATIVE_FONT,
        ),
    )

    worksheet.conditional_formatting.add(
        cell_range,
        CellIsRule(
            operator="equal",
            formula=["0"],
            fill=NEUTRAL_FILL,
            font=NEUTRAL_FONT,
        ),
    )


# ---------------------------------------------------------------------------
# Worksheet layout helpers
# ---------------------------------------------------------------------------

def set_column_widths(
    worksheet: Worksheet,
    widths: dict[int | str, float],
) -> None:
    """Set worksheet column widths."""

    for column, width in widths.items():
        column_letter = (
            get_column_letter(column)
            if isinstance(column, int)
            else column
        )

        worksheet.column_dimensions[column_letter].width = width


def autosize_used_columns(
    worksheet: Worksheet,
    *,
    minimum_width: int = 10,
    maximum_width: int = 35,
    padding: int = 3,
    sample_rows: int = 500,
) -> None:
    """Autosize worksheet columns without requiring a DataFrame."""

    if worksheet.max_column <= 0:
        return

    last_row = min(
        worksheet.max_row,
        sample_rows,
    )

    for column_index in range(1, worksheet.max_column + 1):
        maximum_length = 0

        for row_index in range(1, last_row + 1):
            value = worksheet.cell(
                row=row_index,
                column=column_index,
            ).value

            if value is None:
                continue

            maximum_length = max(
                maximum_length,
                len(str(value)),
            )

        width = max(
            minimum_width,
            min(maximum_length + padding, maximum_width),
        )

        worksheet.column_dimensions[
            get_column_letter(column_index)
        ].width = width


def apply_borders(
    worksheet: Worksheet,
    *,
    start_row: int,
    start_column: int,
    end_row: int,
    end_column: int,
) -> None:
    """Apply the standard report border to a rectangular range."""

    border = thin_border()

    for row in worksheet.iter_rows(
        min_row=start_row,
        max_row=end_row,
        min_col=start_column,
        max_col=end_column,
    ):
        for cell in row:
            cell.border = border


def align_numeric_columns(
    worksheet: Worksheet,
    *,
    columns: Iterable[int | str],
    start_row: int,
    end_row: int,
) -> None:
    """Right-align numeric worksheet columns."""

    for column in columns:
        column_index = (
            column
            if isinstance(column, int)
            else worksheet[column][0].column
        )

        for row in range(start_row, end_row + 1):
            worksheet.cell(
                row=row,
                column=column_index,
            ).alignment = Alignment(
                horizontal="right",
                vertical="center",
            )


def hide_worksheet(
    worksheet: Worksheet,
    *,
    very_hidden: bool = False,
) -> None:
    """Hide a supporting worksheet."""

    worksheet.sheet_state = (
        "veryHidden"
        if very_hidden
        else "hidden"
    )