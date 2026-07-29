"""
Finance Utility Suite
Portfolio Analytics Worksheet Builder

Author  : Shiva Kumar
Version : 1.34.0
"""

from __future__ import annotations

from typing import Any

import pandas as pd
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet

from desktop.reports.excel_styles import (
    THEME,
    centered,
    normal_font,
    section_font,
    solid_fill,
    thin_border,
    title_font,
    value_font,
)


class PortfolioAnalyticsBuilder:
    """Build the Portfolio Analytics worksheet for client holdings reports."""

    CURRENCY_FORMAT = '₹#,##0.00;[Red]-₹#,##0.00'
    PERCENT_FORMAT = '0.00"%"'

    def __init__(
        self,
        account_id: str,
        holdings_df: pd.DataFrame,
        metrics: dict[str, float],
        summary_data: dict[str, Any] | None = None,
    ) -> None:
        self.account_id = str(account_id).strip()
        self.holdings_df = holdings_df.copy()
        self.metrics = dict(metrics)
        self.summary_data = dict(summary_data or {})

    def build(self, worksheet: Worksheet) -> None:
        """Populate and format the complete analytics worksheet."""

        worksheet.sheet_view.showGridLines = False
        self._write_title(worksheet)
        self._write_overview(worksheet)
        self._write_top_holdings(worksheet)
        self._write_mtm_tables(worksheet)
        self._write_margin_table(worksheet)
        self._write_risk_distribution(worksheet)
        self._add_charts(worksheet)
        self._apply_page_setup(worksheet)

    def _write_title(self, worksheet: Worksheet) -> None:
        worksheet.merge_cells("A1:L1")
        worksheet["A1"] = "PORTFOLIO ANALYTICS"
        worksheet["A1"].fill = solid_fill(THEME.dark_blue)
        worksheet["A1"].font = title_font(18)
        worksheet["A1"].alignment = centered()
        worksheet.row_dimensions[1].height = 30

        worksheet.merge_cells("A2:L2")
        worksheet["A2"] = f"Client Account: {self.account_id}"
        worksheet["A2"].fill = solid_fill(THEME.medium_blue)
        worksheet["A2"].font = title_font(11)
        worksheet["A2"].alignment = centered()
        worksheet.row_dimensions[2].height = 22

    def _write_overview(self, worksheet: Worksheet) -> None:
        worksheet.merge_cells("A4:H4")
        worksheet["A4"] = "PORTFOLIO OVERVIEW"
        worksheet["A4"].fill = solid_fill(THEME.dark_blue)
        worksheet["A4"].font = section_font()

        health_score = self._health_score()
        health_label = self._health_label(health_score)
        risk_level = str(self.summary_data.get("Risk Level", "Unknown"))

        cards = [
            ("Total Holdings", self.metrics["holding_count"], "0"),
            ("Total Exposure", self.metrics["total_exposure"], self.CURRENCY_FORMAT),
            ("Total MTM", self.metrics["total_mtm"], self.CURRENCY_FORMAT),
            ("Total Margin", self.metrics["total_margin"], self.CURRENCY_FORMAT),
            ("Largest Holding", self.metrics["largest_holding"] / 100, "0.00%"),
            ("Top 3 Holdings", self.metrics["top_three"] / 100, "0.00%"),
            ("Diversification", self.metrics["diversification"], "0.0"),
            ("Concentration", self.metrics["concentration"], "0.0"),
        ]
        positions = [(6, 1), (6, 3), (6, 5), (6, 7), (9, 1), (9, 3), (9, 5), (9, 7)]
        border = thin_border()

        for (label, value, number_format), (row, column) in zip(cards, positions):
            worksheet.merge_cells(
                start_row=row, start_column=column,
                end_row=row, end_column=column + 1,
            )
            worksheet.merge_cells(
                start_row=row + 1, start_column=column,
                end_row=row + 1, end_column=column + 1,
            )
            label_cell = worksheet.cell(row, column, label)
            value_cell = worksheet.cell(row + 1, column, value)
            label_cell.fill = solid_fill(THEME.pale_blue)
            value_cell.fill = solid_fill(THEME.white)
            label_cell.font = Font(
                name="Segoe UI", size=9, bold=True, color=THEME.muted_text
            )
            value_cell.font = value_font(13)
            label_cell.alignment = centered()
            value_cell.alignment = centered()
            value_cell.number_format = number_format
            for card_row in (row, row + 1):
                for card_column in (column, column + 1):
                    worksheet.cell(card_row, card_column).border = border

        worksheet.merge_cells("J4:L4")
        worksheet["J4"] = "PORTFOLIO HEALTH"
        worksheet["J4"].fill = solid_fill(THEME.dark_blue)
        worksheet["J4"].font = section_font()
        worksheet["J4"].alignment = centered()

        worksheet.merge_cells("J5:L8")
        worksheet["J5"] = health_score
        worksheet["J5"].number_format = '0.0" / 100"'
        worksheet["J5"].font = value_font(24, self._health_color(health_score))
        worksheet["J5"].alignment = centered()
        worksheet["J5"].fill = solid_fill(THEME.pale_blue)

        worksheet.merge_cells("J9:L9")
        worksheet["J9"] = health_label
        worksheet["J9"].font = Font(
            name="Segoe UI", size=12, bold=True, color=THEME.white
        )
        worksheet["J9"].fill = solid_fill(self._health_color(health_score))
        worksheet["J9"].alignment = centered()

        worksheet.merge_cells("J10:L10")
        worksheet["J10"] = f"Risk Level: {risk_level}"
        worksheet["J10"].font = normal_font(10)
        worksheet["J10"].alignment = centered()

    def _write_top_holdings(self, worksheet: Worksheet) -> None:
        start_row = 13
        self._section(worksheet, "A", "D", start_row, "TOP 10 HOLDINGS BY EXPOSURE")
        headers = ["Rank", "Symbol", "Exposure", "Holding %"]
        self._write_headers(worksheet, start_row + 1, 1, headers)

        data = self.holdings_df.sort_values("Exposure", ascending=False).head(10)
        first_data_row = start_row + 2
        for rank, (_, row) in enumerate(data.iterrows(), start=1):
            target = first_data_row + rank - 1
            worksheet.cell(target, 1, rank)
            worksheet.cell(target, 2, str(row.get("Symbol", "")))
            worksheet.cell(target, 3, self._number(row.get("Exposure", 0)))
            worksheet.cell(target, 4, self._number(row.get("Holding %", 0)))
            worksheet.cell(target, 3).number_format = self.CURRENCY_FORMAT
            worksheet.cell(target, 4).number_format = self.PERCENT_FORMAT
            self._style_data_row(worksheet, target, 1, 4, rank)

        self._add_table(worksheet, "TopHoldingsTable", start_row + 1, 1, max(first_data_row + len(data) - 1, start_row + 1), 4)

    def _write_mtm_tables(self, worksheet: Worksheet) -> None:
        start_row = 13
        self._section(worksheet, "F", "H", start_row, "TOP MTM GAINERS")
        self._write_headers(worksheet, start_row + 1, 6, ["Rank", "Symbol", "MTM"])
        gainers = self.holdings_df[self.holdings_df["MarkToMarket"] > 0].nlargest(10, "MarkToMarket")
        self._write_ranked_value_table(worksheet, gainers, start_row + 2, 6, "MarkToMarket")

        losers_start = 27
        self._section(worksheet, "F", "H", losers_start, "TOP MTM LOSERS")
        self._write_headers(worksheet, losers_start + 1, 6, ["Rank", "Symbol", "MTM"])
        losers = self.holdings_df[self.holdings_df["MarkToMarket"] < 0].nsmallest(10, "MarkToMarket")
        self._write_ranked_value_table(worksheet, losers, losers_start + 2, 6, "MarkToMarket")

    def _write_margin_table(self, worksheet: Worksheet) -> None:
        start_row = 27
        self._section(worksheet, "A", "D", start_row, "TOP MTF MARGIN STOCKS")
        self._write_headers(worksheet, start_row + 1, 1, ["Rank", "Symbol", "MTF Margin", "Exposure"])
        data = self.holdings_df.nlargest(10, "MTF MARGIN")
        for rank, (_, row) in enumerate(data.iterrows(), start=1):
            target = start_row + 1 + rank
            values = [rank, str(row.get("Symbol", "")), self._number(row.get("MTF MARGIN", 0)), self._number(row.get("Exposure", 0))]
            for offset, value in enumerate(values):
                worksheet.cell(target, 1 + offset, value)
            worksheet.cell(target, 3).number_format = self.CURRENCY_FORMAT
            worksheet.cell(target, 4).number_format = self.CURRENCY_FORMAT
            self._style_data_row(worksheet, target, 1, 4, rank)

    def _write_risk_distribution(self, worksheet: Worksheet) -> None:
        start_row = 42
        self._section(worksheet, "A", "D", start_row, "CONCENTRATION RISK DISTRIBUTION")
        self._write_headers(worksheet, start_row + 1, 1, ["Risk Band", "Holdings", "Exposure", "Exposure %"])

        classified = self.holdings_df.copy()
        classified["Risk Band"] = classified["Holding %"].apply(self._risk_band)
        total_exposure = max(float(classified["Exposure"].sum()), 0.0)
        rows = []
        for band in ["Low", "Medium", "High", "Critical"]:
            subset = classified[classified["Risk Band"] == band]
            exposure = float(subset["Exposure"].sum())
            rows.append((band, len(subset), exposure, (exposure / total_exposure * 100) if total_exposure else 0.0))

        for index, values in enumerate(rows, start=1):
            target = start_row + 1 + index
            for offset, value in enumerate(values):
                worksheet.cell(target, 1 + offset, value)
            worksheet.cell(target, 3).number_format = self.CURRENCY_FORMAT
            worksheet.cell(target, 4).number_format = self.PERCENT_FORMAT
            worksheet.cell(target, 1).font = Font(name="Segoe UI", size=10, bold=True, color=self._risk_color(values[0]))
            self._style_data_row(worksheet, target, 1, 4, index)

    def _add_charts(self, worksheet: Worksheet) -> None:
        # Top holdings chart from A14:D24.
        top_end = min(24, 14 + min(10, len(self.holdings_df)))
        if top_end >= 15:
            chart = BarChart()
            chart.type = "bar"
            chart.style = 10
            chart.title = "Top Holdings by Exposure"
            chart.height = 7.6
            chart.width = 12.8
            chart.legend = None
            data = Reference(worksheet, min_col=3, min_row=14, max_row=top_end)
            cats = Reference(worksheet, min_col=2, min_row=15, max_row=top_end)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True
            worksheet.add_chart(chart, "J13")

        # Risk distribution pie chart.
        chart = PieChart()
        chart.title = "Risk Distribution by Exposure"
        chart.height = 7.2
        chart.width = 12.0
        data = Reference(worksheet, min_col=3, min_row=43, max_row=47)
        labels = Reference(worksheet, min_col=1, min_row=44, max_row=47)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(labels)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showPercent = True
        chart.legend.position = "r"
        worksheet.add_chart(chart, "F42")

    def _write_ranked_value_table(
        self,
        worksheet: Worksheet,
        dataframe: pd.DataFrame,
        start_row: int,
        start_column: int,
        value_column: str,
    ) -> None:
        for rank, (_, row) in enumerate(dataframe.iterrows(), start=1):
            target = start_row + rank - 1
            worksheet.cell(target, start_column, rank)
            worksheet.cell(target, start_column + 1, str(row.get("Symbol", "")))
            worksheet.cell(target, start_column + 2, self._number(row.get(value_column, 0)))
            worksheet.cell(target, start_column + 2).number_format = self.CURRENCY_FORMAT
            self._style_data_row(worksheet, target, start_column, start_column + 2, rank)

    @staticmethod
    def _section(worksheet: Worksheet, start_col: str, end_col: str, row: int, text: str) -> None:
        worksheet.merge_cells(f"{start_col}{row}:{end_col}{row}")
        cell = worksheet[f"{start_col}{row}"]
        cell.value = text
        cell.fill = solid_fill(THEME.dark_blue)
        cell.font = section_font()
        cell.alignment = Alignment(horizontal="left", vertical="center")

    @staticmethod
    def _write_headers(worksheet: Worksheet, row: int, start_column: int, headers: list[str]) -> None:
        for offset, header in enumerate(headers):
            cell = worksheet.cell(row, start_column + offset, header)
            cell.fill = solid_fill(THEME.medium_blue)
            cell.font = Font(name="Segoe UI", size=9, bold=True, color=THEME.white)
            cell.alignment = centered()
            cell.border = thin_border()

    @staticmethod
    def _style_data_row(worksheet: Worksheet, row: int, start_column: int, end_column: int, index: int) -> None:
        fill = solid_fill("F8FAFC" if index % 2 == 0 else THEME.white)
        for column in range(start_column, end_column + 1):
            cell = worksheet.cell(row, column)
            cell.fill = fill
            cell.border = thin_border()
            cell.font = normal_font(9)
            cell.alignment = Alignment(horizontal="center" if column == start_column else "right", vertical="center")

    @staticmethod
    def _add_table(worksheet: Worksheet, name: str, min_row: int, min_col: int, max_row: int, max_col: int) -> None:
        if max_row <= min_row:
            return
        start = worksheet.cell(min_row, min_col).coordinate
        end = worksheet.cell(max_row, max_col).coordinate
        table = Table(displayName=name, ref=f"{start}:{end}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    def _health_score(self) -> float:
        diversification = self.metrics.get("diversification", 0.0)
        concentration_quality = 100.0 - self.metrics.get("concentration", 100.0)
        largest = self.metrics.get("largest_holding", 100.0)
        largest_quality = max(0.0, min(100.0, 100.0 - (largest * 2.5)))
        exposure = abs(self.metrics.get("total_exposure", 0.0))
        mtm = self.metrics.get("total_mtm", 0.0)
        mtm_quality = 50.0 if exposure <= 0 else max(0.0, min(100.0, 50.0 + (mtm / exposure * 250.0)))
        return max(0.0, min(100.0, diversification * 0.40 + concentration_quality * 0.20 + largest_quality * 0.20 + mtm_quality * 0.20))

    @staticmethod
    def _health_label(score: float) -> str:
        if score >= 80:
            return "Excellent"
        if score >= 65:
            return "Good"
        if score >= 50:
            return "Moderate"
        if score >= 35:
            return "Weak"
        return "Critical"

    @staticmethod
    def _health_color(score: float) -> str:
        if score >= 80:
            return THEME.green
        if score >= 65:
            return "65A30D"
        if score >= 50:
            return THEME.amber
        if score >= 35:
            return THEME.orange
        return THEME.red

    @staticmethod
    def _risk_band(value: Any) -> str:
        percentage = PortfolioAnalyticsBuilder._number(value)
        if percentage > 35:
            return "Critical"
        if percentage > 20:
            return "High"
        if percentage >= 10:
            return "Medium"
        return "Low"

    @staticmethod
    def _risk_color(band: str) -> str:
        return {
            "Low": THEME.green,
            "Medium": THEME.amber,
            "High": THEME.orange,
            "Critical": THEME.red,
        }.get(band, THEME.muted_text)

    @staticmethod
    def _number(value: Any) -> float:
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    @staticmethod
    def _apply_page_setup(worksheet: Worksheet) -> None:
        for letter, width in {
            "A": 11, "B": 17, "C": 18, "D": 14,
            "E": 3, "F": 11, "G": 17, "H": 18,
            "I": 3, "J": 16, "K": 16, "L": 16,
        }.items():
            worksheet.column_dimensions[letter].width = width
        worksheet.freeze_panes = "A4"
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.print_area = "A1:L50"
        worksheet.oddFooter.center.text = "Finance Utility Suite | Confidential | Page &P of &N"
