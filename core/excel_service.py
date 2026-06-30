"""
Excel Service
Finance Utility Suite

Creates professional Excel reports.
"""

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from config import OUTPUT_FOLDER
from core.logger import AppLogger


class ExcelService:

    def __init__(self):

        self.logger = AppLogger()

    # ----------------------------------------
    # Export DataFrame
    # ----------------------------------------

    def export(self, dataframe):

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        filename = OUTPUT_FOLDER / f"Stock_Report_{timestamp}.xlsx"

        wb = Workbook()

        ws = wb.active

        ws.title = "Stock Analysis"

        # ----------------------------------------
        # Styles
        # ----------------------------------------

        header_fill = PatternFill(
            fill_type="solid",
            start_color="1F4E78",
            end_color="1F4E78"
        )

        header_font = Font(
            bold=True,
            color="FFFFFF"
        )

        center = Alignment(horizontal="center")

        # ----------------------------------------
        # Header
        # ----------------------------------------

        for col, column in enumerate(dataframe.columns, start=1):

            cell = ws.cell(row=1, column=col)

            cell.value = column

            cell.fill = header_fill

            cell.font = header_font

            cell.alignment = center

        # ----------------------------------------
        # Data
        # ----------------------------------------

        for row in dataframe.itertuples(index=False):

            ws.append(list(row))

        # ----------------------------------------
        # Auto Width
        # ----------------------------------------

        for column in ws.columns:

            max_length = 0

            column_letter = get_column_letter(column[0].column)

            for cell in column:

                try:

                    if len(str(cell.value)) > max_length:

                        max_length = len(str(cell.value))

                except:

                    pass

            ws.column_dimensions[column_letter].width = max_length + 3

        # ----------------------------------------
        # Freeze Header
        # ----------------------------------------

        ws.freeze_panes = "A2"

        # ----------------------------------------
        # Filter
        # ----------------------------------------

        ws.auto_filter.ref = ws.dimensions

        # ----------------------------------------
        # Summary Sheet
        # ----------------------------------------

        summary = wb.create_sheet("Summary")

        summary["A1"] = "Metric"
        summary["B1"] = "Value"

        summary["A1"].font = header_font
        summary["B1"].font = header_font

        summary["A1"].fill = header_fill
        summary["B1"].fill = header_fill

        summary.append(["Total Stocks", len(dataframe)])

        if "Risk" in dataframe.columns:

            summary.append([
                "High Risk",
                len(dataframe[dataframe["Risk"] == "High"])
            ])

            summary.append([
                "Moderate Risk",
                len(dataframe[dataframe["Risk"] == "Moderate"])
            ])

            summary.append([
                "Low Risk",
                len(dataframe[dataframe["Risk"] == "Low"])
            ])

        if "Beta" in dataframe.columns:

            summary.append([
                "Average Beta",
                round(dataframe["Beta"].mean(), 2)
            ])

        wb.save(filename)

        self.logger.info(f"Excel exported : {filename}")

        return filename