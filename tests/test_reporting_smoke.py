from __future__ import annotations
import tempfile, unittest
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from desktop.reports.portfolio_report import PortfolioReport
from desktop.reports.report_controller import ReportController
from desktop.reports.report_engine import ReportEngineError
class PortfolioReportTests(unittest.TestCase):
    @staticmethod
    def sample_data():
        return pd.DataFrame([{"Symbol":"RELIANCE","Quantity":10,"Average Price":2500,"Current Price":2700,"Sector":"Energy"},{"Symbol":"TCS","Quantity":5,"Average Price":3500,"Current Price":3900,"Sector":"Technology"},{"Symbol":"RELIANCE","Quantity":2,"Average Price":2600,"Current Price":2700,"Sector":"Energy"}])
    def test_portfolio_export(self):
        with tempfile.TemporaryDirectory() as d:
            result=ReportController.generate_portfolio(Path(d)/"portfolio.xlsx", portfolio_df=self.sample_data(), portfolio_name="Test Portfolio")
            wb=load_workbook(result); self.assertEqual(wb.sheetnames[:4],["Dashboard","Holdings","Performance","Allocation"]); self.assertEqual(wb["Chart Data"].sheet_state,"hidden"); self.assertEqual(wb.active.title,"Dashboard"); wb.close()
    def test_duplicate_symbols_are_consolidated(self):
        prepared=PortfolioReport(self.sample_data())._prepare_dataframe(self.sample_data()); self.assertEqual(len(prepared),2); self.assertEqual(float(prepared.loc[prepared["Symbol"]=="RELIANCE","Quantity"].iloc[0]),12.0)
    def test_empty_portfolio_fails_cleanly(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(ReportEngineError) as c: PortfolioReport(pd.DataFrame()).export(Path(d)/"empty.xlsx")
            self.assertIn("Portfolio data is empty",str(c.exception))
    def test_missing_columns_fail_cleanly(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(ReportEngineError) as c: PortfolioReport(pd.DataFrame([{"Symbol":"TCS"}])).export(Path(d)/"missing.xlsx")
            self.assertIn("Missing required portfolio columns",str(c.exception))
if __name__ == "__main__": unittest.main()
