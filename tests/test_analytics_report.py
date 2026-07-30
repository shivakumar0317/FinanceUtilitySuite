from __future__ import annotations
import tempfile
import unittest
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from desktop.reports.analytics_engine import AnalyticsDataError, PortfolioAnalyticsEngine
from desktop.reports.report_controller import ReportController
from desktop.reports.report_engine import ReportEngineError

class AnalyticsReportTests(unittest.TestCase):
    @staticmethod
    def data():
        return pd.DataFrame([
            {"Symbol":"RELIANCE","Quantity":10,"Average Price":2500,"Current Price":2700,"Sector":"Energy"},
            {"Symbol":"TCS","Quantity":5,"Average Price":3500,"Current Price":3900,"Sector":"Technology"},
            {"Symbol":"RELIANCE","Quantity":2,"Average Price":2600,"Current Price":2700,"Sector":"Energy"},
        ])
    def test_kpi_calculation(self):
        s=PortfolioAnalyticsEngine().summarize(self.data())
        self.assertEqual(s.holdings_count,2)
        self.assertEqual(s.sector_count,2)
        self.assertEqual(s.total_investment,47700.0)
        self.assertEqual(s.current_value,51900.0)
        self.assertEqual(s.profit_loss,4200.0)
        self.assertEqual(s.largest_holding_symbol,"RELIANCE")
        self.assertGreaterEqual(s.portfolio_health_score,0)
        self.assertLessEqual(s.portfolio_health_score,100)
    def test_duplicate_symbols_are_consolidated(self):
        frame=PortfolioAnalyticsEngine().prepare(self.data())
        self.assertEqual(len(frame),2)
        self.assertEqual(float(frame.loc[frame.Symbol=="RELIANCE","Quantity"].iloc[0]),12.0)
    def test_empty_data(self):
        with self.assertRaises(AnalyticsDataError): PortfolioAnalyticsEngine().summarize(pd.DataFrame())
    def test_missing_columns(self):
        with self.assertRaises(AnalyticsDataError): PortfolioAnalyticsEngine().summarize(pd.DataFrame([{"Symbol":"TCS"}]))
    def test_workbook_generation(self):
        with tempfile.TemporaryDirectory() as folder:
            output=ReportController.generate_analytics(Path(folder)/"analytics.xlsx", portfolio_df=self.data(), portfolio_name="Test")
            self.assertTrue(output.exists())
            wb=load_workbook(output)
            self.assertEqual(wb.sheetnames,["Dashboard","Performance Summary","Monthly Returns","Sector Analysis","Top Performers","Summary","Chart Data"])
            self.assertEqual(wb["Chart Data"].sheet_state,"hidden")
            self.assertEqual(wb.active.title,"Dashboard")
            self.assertGreaterEqual(len(wb["Dashboard"]._charts),2)
            wb.close()

if __name__=='__main__': unittest.main()
