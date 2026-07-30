from __future__ import annotations

import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd
from openpyxl import load_workbook

from desktop.reports.portfolio_report import PortfolioReport


class PortfolioWorkbookStabilityTests(unittest.TestCase):

    def setUp(self) -> None:
        self.dataframe = pd.DataFrame(
            {
                "Symbol": ["RELIANCE", "TCS", "INFY"],
                "Quantity": [10, 5, 7],
                "Average Price": [100.0, 200.0, 300.0],
                "Current Price": [110.0, 190.0, 330.0],
                "Sector": ["Energy", "Technology", "Technology"],
            }
        )

    def test_generated_workbook_reopens_and_contains_expected_tables(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "portfolio_report.xlsx"
            PortfolioReport(
                self.dataframe,
                portfolio_name="Stability Test",
            ).export(output)

            workbook = load_workbook(output)
            self.assertEqual(
                workbook.sheetnames,
                ["Dashboard", "Holdings", "Performance", "Allocation", "Chart Data"],
            )
            self.assertEqual(
                list(workbook["Holdings"].tables),
                ["PortfolioHoldings"],
            )
            self.assertEqual(
                list(workbook["Performance"].tables),
                ["PortfolioPerformance"],
            )
            self.assertEqual(
                set(workbook["Allocation"].tables),
                {"SectorAllocation", "HoldingAllocation"},
            )

    def test_all_table_column_ids_are_sequential(self) -> None:
        namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "portfolio_report.xlsx"
            PortfolioReport(
                self.dataframe,
                portfolio_name="Stability Test",
            ).export(output)

            with zipfile.ZipFile(output) as archive:
                table_files = sorted(
                    name for name in archive.namelist()
                    if name.startswith("xl/tables/table") and name.endswith(".xml")
                )
                self.assertEqual(len(table_files), 4)

                for table_file in table_files:
                    root = ET.fromstring(archive.read(table_file))
                    columns = root.findall("x:tableColumns/x:tableColumn", namespace)
                    ids = [int(column.attrib["id"]) for column in columns]
                    self.assertEqual(ids, list(range(1, len(columns) + 1)))


if __name__ == "__main__":
    unittest.main()
