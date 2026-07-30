from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from desktop.reports.analytics_report import AnalyticsReport
from desktop.reports.report_controller import ReportController


class AnalyticsPhase2Tests(unittest.TestCase):
    @staticmethod
    def holdings() -> pd.DataFrame:
        return pd.DataFrame([
            {"Symbol":"ALPHA","Quantity":10,"Average Price":100,"Current Price":130,"Sector":"Technology"},
            {"Symbol":"BETA","Quantity":20,"Average Price":50,"Current Price":45,"Sector":"Finance"},
            {"Symbol":"GAMMA","Quantity":5,"Average Price":200,"Current Price":220,"Sector":"Technology"},
        ])

    @staticmethod
    def history() -> pd.DataFrame:
        return pd.DataFrame({
            "Date":["2024-01-01","2024-01-31","2024-02-29","2024-03-31","2025-01-31"],
            "Portfolio Value":[100000,105000,110000,108000,121000],
        })

    def export(self, with_history: bool = True):
        tmp=tempfile.TemporaryDirectory()
        path=Path(tmp.name)/"analytics.xlsx"
        AnalyticsReport(self.holdings(), portfolio_name="Test", history_dataframe=self.history() if with_history else None).export(path)
        return tmp,path,load_workbook(path)

    def test_phase2_sheet_structure(self):
        tmp,path,wb=self.export()
        self.addCleanup(tmp.cleanup)
        self.assertEqual(wb.sheetnames,["Dashboard","Performance Summary","Monthly Returns","Sector Analysis","Top Performers","Summary","Chart Data"])
        self.assertEqual(wb.active.title,"Dashboard")
        self.assertEqual(wb["Chart Data"].sheet_state,"hidden")

    def test_dashboard_contains_two_charts(self):
        tmp,path,wb=self.export()
        self.addCleanup(tmp.cleanup)
        self.assertGreaterEqual(len(wb["Dashboard"]._charts),2)

    def test_history_populates_returns_and_growth_charts(self):
        tmp,path,wb=self.export()
        self.addCleanup(tmp.cleanup)
        self.assertGreater(wb["Monthly Returns"].max_row,5)
        self.assertGreaterEqual(len(wb["Monthly Returns"]._charts),1)
        self.assertGreaterEqual(len(wb["Performance Summary"]._charts),1)

    def test_without_history_exports_cleanly(self):
        tmp,path,wb=self.export(False)
        self.addCleanup(tmp.cleanup)
        self.assertIn("Historical portfolio data was not supplied",wb["Monthly Returns"]["A4"].value)
        self.assertEqual(len(wb["Monthly Returns"]._charts),0)

    def test_sector_and_performer_tables(self):
        tmp,path,wb=self.export()
        self.addCleanup(tmp.cleanup)
        self.assertIn("SectorAnalysisTable",wb["Sector Analysis"].tables)
        self.assertIn("TopWinnersTable",wb["Top Performers"].tables)
        self.assertIn("TopLosersTable",wb["Top Performers"].tables)

    def test_controller_accepts_history(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/"controller.xlsx"
            result=ReportController.generate_analytics(path,portfolio_df=self.holdings(),portfolio_name="Controller",history_df=self.history())
            self.assertTrue(result.exists())


if __name__ == "__main__":
    unittest.main()
