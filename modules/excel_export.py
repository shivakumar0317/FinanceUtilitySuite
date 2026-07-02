from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


class ExcelExporter:

    def export(self, dataframe, filename):

        wb = Workbook()

        ws = wb.active
        ws.title = "Stock Analysis"

        header_fill = PatternFill(
            start_color="1F4E78", end_color="1F4E78", fill_type="solid"
        )

        header_font = Font(bold=True, color="FFFFFF")

        # Write Header
        for col, column in enumerate(dataframe.columns, start=1):

            cell = ws.cell(row=1, column=col)

            cell.value = column
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        # Write Data
        for row in dataframe.itertuples(index=False):
            ws.append(list(row))

        # Auto Width
        for column in ws.columns:

            length = max(len(str(cell.value)) if cell.value else 0 for cell in column)

            ws.column_dimensions[get_column_letter(column[0].column)].width = length + 3

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        wb.save(filename)
